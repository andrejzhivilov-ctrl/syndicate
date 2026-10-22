"""Fin — звіти та експорт: Markdown, XLSX, DOCX, PDF, PPTX, Obsidian.

Звіти збираються з чистих даних (fin_ledger). Кожен звіт — повний або короткий для Єви —
ЗАВЖДИ містить блок «Заплановано» (день / тиждень / місяць / наступні 30 днів) з сумами.

Google Sheets / Docs / Drive тут НЕ викликаються: це окремі дозволи (fin.sheets.cockpit,
fin.docs.generate, fin.drive.archive). Без дозволу повертається blocked_not_granted.
"""

from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Dict, List, Optional

from src.agents import fin_ledger as L

EXPORT_FORMATS = ("md", "xlsx", "docx", "pdf", "pptx")


def _fmt_totals(totals: Dict[str, str]) -> str:
    if not totals:
        return "0"
    return " + ".join(f"{v} {c}" for c, v in totals.items())


def build_report(kind: str, today: date, payments: List[L.PlannedPayment],
                 expenses: List[L.Expense], fund_balances: Dict[str, Decimal]) -> dict:
    """kind: full (повний, ЗВІТ) | standard (день/тиждень/місяць) | eva (короткий для Єви)."""
    if kind == "eva":
        period = ("day", *L.period_bounds(today, "day"))
    elif kind in ("day", "week", "month"):
        period = (kind, *L.period_bounds(today, kind))
    else:
        period = ("month", *L.period_bounds(today, "month"))
    pname, ps, pe = period[0], period[1], period[2]

    header = L.plan_header(payments, today)
    items = L.plan_items(payments, ps, pe)
    spent = L.spend_by_fund(expenses, ps, pe)
    spent_rows = [e for e in expenses if ps <= e.on <= pe]
    planned_fund = {f: Decimal("0") for f in L.FUNDS}
    for it in items:
        planned_fund[it["fund"]] += Decimal(it["amount"]) if it["currency"] == "UAH" else Decimal("0")

    funds = []
    for f, label in L.FUNDS.items():
        bal = Decimal(str(fund_balances.get(f, 0)))
        funds.append({"id": f, "label": label, "balance": str(bal),
                      "planned_in_period": str(planned_fund[f]),
                      "spent_in_period": str(spent[f]),
                      "available": str(L.available_in_fund(bal, planned_fund[f]))})

    report = {
        "kind": kind,
        "title": {"full": "ФІН — повний фінансовий звіт", "standard": "ФІН — стандартний звіт",
                  "eva": "ФІН — для Єви"}.get(kind, "ФІН — звіт"),
        "generated_at": datetime.now().isoformat(timespec="minutes"),
        "period": {"name": pname, "from": ps.isoformat(), "to": pe.isoformat()},
        "planned_header": header,             # ЗАВЖДИ: що заплановано і суми
        "planned_items": items,
        "funds": funds,
        "expenses": [{"date": e.on.isoformat(), "amount": str(e.amount), "currency": e.currency,
                      "fund": e.fund, "category": e.category, "store": e.store,
                      "purpose": e.purpose, "source": e.source, "ordered_by": e.ordered_by}
                     for e in spent_rows],
        "spent_totals": {"by_fund": {k: str(v) for k, v in spent.items()},
                         "total": str(sum(spent.values(), Decimal("0")))},
        "sources": [] if kind == "eva" else [
            {"id": sid, "label": m["label"], "cash_path": m["cash_path"], "intl": m["intl"]}
            for sid, m in L.SOURCES.items()],
        "notes": ["Фін лише спостерігає: платежі й замовлення робить Андрій.",
                  "Фонд безпеки не є джерелом покупок.",
                  "Суми з різних валют не складаються між собою."],
    }
    return report


# ---------------------------------------------------------------- Markdown
def to_markdown(r: dict) -> str:
    h = r["planned_header"]
    lines = [f"# {r['title']}", f"_Період: {r['period']['from']} — {r['period']['to']}, "
             f"згенеровано {r['generated_at']}_", "",
             "## Заплановано (завжди)",
             "| Період | К-сть | Сума |", "|---|---|---|"]
    for key, label in (("day", "День"), ("week", "Тиждень"), ("month", "Місяць"),
                       ("next_30_days", "Наступні 30 днів")):
        lines.append(f"| {label} | {h[key]['count']} | {_fmt_totals(h[key]['totals'])} |")
    lines += ["", "## Фонди", "| Фонд | Залишок | Заплановано | Витрачено | Доступно |",
              "|---|---|---|---|---|"]
    for f in r["funds"]:
        lines.append(f"| {f['label']} | {f['balance']} | {f['planned_in_period']} | "
                     f"{f['spent_in_period']} | {f['available']} |")
    lines += ["", "## Плановані платежі (період)"]
    lines += [f"- {i['date']} — {i['name']}: {i['amount']} {i['currency']} ({i['fund']})"
              for i in r["planned_items"]] or ["- немає"]
    lines += ["", "## Витрати (період)"]
    lines += [f"- {e['date']} — {e['purpose']} ({e['store']}): {e['amount']} {e['currency']} "
              f"→ {e['fund']}, {e['source']}" for e in r["expenses"]] or ["- немає"]
    if r["sources"]:
        lines += ["", "## Джерела коштів (без реквізитів)"]
        lines += [f"- {s['label']} — {s['cash_path']}{' (міжнародний)' if s['intl'] else ''}"
                  for s in r["sources"]]
    lines += ["", "## Примітки"] + [f"- {n}" for n in r["notes"]]
    return "\n".join(lines) + "\n"


def _rows_for(r: dict) -> List[List[str]]:
    """Плоскі рядки для табличних форматів (XLSX, DOCX, PPTX, PDF)."""
    h = r["planned_header"]
    rows = [["Заплановано", "Кількість", "Сума"]]
    for key, label in (("day", "День"), ("week", "Тиждень"), ("month", "Місяць"),
                       ("next_30_days", "Наступні 30 днів")):
        rows.append([label, str(h[key]["count"]), _fmt_totals(h[key]["totals"])])
    return rows


def _funds_rows(r: dict) -> List[List[str]]:
    rows = [["Фонд", "Залишок", "Заплановано", "Витрачено", "Доступно"]]
    rows += [[f["label"], f["balance"], f["planned_in_period"], f["spent_in_period"],
              f["available"]] for f in r["funds"]]
    return rows


def _planned_rows(r: dict) -> List[List[str]]:
    rows = [["Дата", "Платіж", "Сума", "Валюта", "Фонд"]]
    rows += [[i["date"], i["name"], i["amount"], i["currency"], i["fund"]] for i in r["planned_items"]]
    return rows


def _expense_rows(r: dict) -> List[List[str]]:
    rows = [["Дата", "Призначення", "Магазин", "Сума", "Валюта", "Фонд", "Джерело"]]
    rows += [[e["date"], e["purpose"], e["store"], e["amount"], e["currency"], e["fund"], e["source"]]
             for e in r["expenses"]]
    return rows


# ---------------------------------------------------------------- XLSX
def write_xlsx(r: dict, path: Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    sheets = [("DASHBOARD", _rows_for(r) + [[]] + _funds_rows(r)),
              ("ПЛАН", _planned_rows(r)),
              ("ВИТРАТИ", _expense_rows(r))]
    for idx, (name, rows) in enumerate(sheets):
        ws = wb.active if idx == 0 else wb.create_sheet()
        ws.title = name
        for row in rows:
            ws.append(row)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = min(
                40, max(12, max(len(str(c.value or "")) for c in col) + 2))
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


# ---------------------------------------------------------------- DOCX
def write_docx(r: dict, path: Path) -> Path:
    from docx import Document

    doc = Document()
    doc.add_heading(r["title"], level=0)
    doc.add_paragraph(f"Період: {r['period']['from']} — {r['period']['to']}")
    for heading, rows in (("Заплановано (завжди)", _rows_for(r)), ("Фонди", _funds_rows(r)),
                          ("Плановані платежі", _planned_rows(r)), ("Витрати", _expense_rows(r))):
        doc.add_heading(heading, level=1)
        table = doc.add_table(rows=0, cols=len(rows[0]))
        table.style = "Table Grid"
        for row in rows:
            cells = table.add_row().cells
            for c, v in zip(cells, row, strict=True):
                c.text = str(v)
    doc.add_heading("Примітки", level=1)
    for n in r["notes"]:
        doc.add_paragraph(n, style="List Bullet")
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    return path


# ---------------------------------------------------------------- PPTX
def write_pptx(r: dict, path: Path) -> Path:
    from pptx import Presentation
    from pptx.util import Pt

    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = r["title"]
    s.placeholders[1].text = f"{r['period']['from']} — {r['period']['to']}"
    for heading, rows in (("Заплановано (завжди)", _rows_for(r)), ("Фонди", _funds_rows(r))):
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = heading
        shape = slide.shapes.add_table(len(rows), len(rows[0]), Pt(30), Pt(120),
                                       prs.slide_width - Pt(60), Pt(30) * len(rows))
        for i, row in enumerate(rows):
            for j, v in enumerate(row):
                shape.table.cell(i, j).text = str(v)
    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(path)
    return path


# ---------------------------------------------------------------- PDF
def write_pdf(r: dict, path: Path) -> Path:
    from PIL import ImageFont
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    # Pillow searches the platform's font directories, including non-Debian layouts.
    regular = bold = None
    for reg_cand, bold_cand in (
        ("DejaVuSans.ttf", "DejaVuSans-Bold.ttf"),
        ("arial.ttf", "arialbd.ttf"),
        ("segoeui.ttf", "segoeuib.ttf"),
    ):
        try:
            regular = ImageFont.truetype(reg_cand, 12).path
            bold = ImageFont.truetype(bold_cand, 12).path
            break
        except OSError:
            continue
    if not regular or not bold:
        raise RuntimeError("PDF: немає шрифту з кирилицею (DejaVuSans)")
    pdfmetrics.registerFont(TTFont("DejaVu", regular))
    pdfmetrics.registerFont(TTFont("DejaVu-Bold", bold))
    st = getSampleStyleSheet()
    st["Title"].fontName = "DejaVu-Bold"
    st["Normal"].fontName = "DejaVu"
    st["Heading2"].fontName = "DejaVu-Bold"

    story = [Paragraph(r["title"], st["Title"]),
             Paragraph(f"Період: {r['period']['from']} — {r['period']['to']}", st["Normal"]),
             Spacer(1, 6 * mm)]
    for heading, rows in (("Заплановано (завжди)", _rows_for(r)), ("Фонди", _funds_rows(r)),
                          ("Плановані платежі", _planned_rows(r)), ("Витрати", _expense_rows(r))):
        story.append(Paragraph(heading, st["Heading2"]))
        t = Table(rows, repeatRows=1)
        t.setStyle(TableStyle([("FONTNAME", (0, 0), (-1, -1), "DejaVu"),
                               ("FONTSIZE", (0, 0), (-1, -1), 8),
                               ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
                               ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e8eef7"))]))
        story += [t, Spacer(1, 4 * mm)]
    story.append(Paragraph("Примітки", st["Heading2"]))
    story += [Paragraph(f"• {n}", st["Normal"]) for n in r["notes"]]
    path.parent.mkdir(parents=True, exist_ok=True)
    SimpleDocTemplate(str(path), pagesize=A4, leftMargin=14 * mm, rightMargin=14 * mm,
                      topMargin=14 * mm, bottomMargin=14 * mm).build(story)
    return path


# ---------------------------------------------------------------- Obsidian / сервіси
def write_obsidian(r: dict, vault_dir: Optional[Path]) -> Path:
    """Пише .md у локальний vault Фіна. Без налаштованого vault — відмова."""
    if vault_dir is None:
        raise PermissionError("obsidian_vault_not_configured")
    folder = Path(vault_dir) / "Fin" / "Звіти"
    folder.mkdir(parents=True, exist_ok=True)
    path = folder / f"{r['kind']}_{r['period']['from']}.md"
    path.write_text(to_markdown(r), encoding="utf-8")
    return path


EXTERNAL_SINKS = {
    "google_sheets": "fin.sheets.cockpit",
    "google_docs": "fin.docs.generate",
    "google_drive": "fin.drive.archive",
}


def external_sink_status(sink: str) -> dict:
    """Зовнішні сервіси — лише статус. Без дозволу нічого не відправляється."""
    cap = EXTERNAL_SINKS.get(sink)
    if cap is None:
        return {"status": "unknown_sink"}
    return {"status": "blocked_not_granted", "needs": cap}


def export(r: dict, fmt: str, out_dir: Path, vault_dir: Optional[Path] = None) -> Path:
    stem = f"fin_{r['kind']}_{r['period']['from']}"
    if fmt == "md":
        out_dir.mkdir(parents=True, exist_ok=True)
        p = out_dir / f"{stem}.md"
        p.write_text(to_markdown(r), encoding="utf-8")
        return p
    if fmt == "obsidian":
        return write_obsidian(r, vault_dir)
    writers = {"xlsx": write_xlsx, "docx": write_docx, "pdf": write_pdf, "pptx": write_pptx}
    if fmt not in writers:
        raise ValueError(f"unsupported format: {fmt}")
    return writers[fmt](r, out_dir / f"{stem}.{fmt}")
